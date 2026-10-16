#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
07_docx_corpus.py —— 真实公文语料提取器

作用
----
把用户提供的 10 份真实公文（两组：建筑垃圾规划汇报 3 份、实施层面详细规划
编制工作流程 7 份迭代稿）提取为纯文本语料，并输出「语料画像」（字符数、段落
数、估算 token 数），供后续「真实场景」性能/仿真测试直接引用。

设计要点
--------
1. 只读：绝不修改源 docx（办公文档为原始证据，保持原样）。
2. 文本不进 Agent 上下文：脚本自行读取并落盘为 corpus/corpus.json，
   控制台只打印画像摘要，避免污染会话上下文预算。
3. 去噪：跳过空段落；把连续空白归一；保留段落自然边界（用 \n 连接），
   因为后续测试要按「段落」切分做分段改写等真实操作。

语料位置
--------
默认读取项目内 `测试数据/真实公文语料/`（两组子目录），随项目迁移、可直接冷启动；
需要改用其它目录时设环境变量 RKLLM_CORPUS_DIR 指向该目录（其下仍按两组子目录组织）。

用法
----
    python 07_docx_corpus.py                 # 提取全部并打印画像
    python 07_docx_corpus.py --out corpus/corpus.json
    RKLLM_CORPUS_DIR=/path/to/语料 python 07_docx_corpus.py   # 用外部语料目录
"""

import argparse
import json
import os
import re
import sys
import time

try:
    from docx import Document
except ImportError:
    sys.stderr.write("缺少 python-docx，请先安装：pip install python-docx\n")
    raise

# ---------------------------------------------------------------------------
# 语料清单：两组真实文档
#   group  —— 业务分组
#   stage  —— 在真实工作流中的阶段（草稿/润色/审核/终稿…），用于还原真实迭代
# ---------------------------------------------------------------------------
# 语料根目录（软编码）：默认为项目内「测试数据/真实公文语料」，随项目一起迁移；
# 需要指向其它位置（如本机原始稿件目录）时设环境变量 RKLLM_CORPUS_DIR 覆盖。
_HERE = os.path.dirname(os.path.abspath(__file__))
CORPUS_ROOT = os.environ.get("RKLLM_CORPUS_DIR") or os.path.normpath(
    os.path.join(_HERE, "..", "..", "测试数据", "真实公文语料"))
BASE_BUILD = os.path.join(CORPUS_ROOT, "建筑垃圾规划布置汇报")
BASE_PLAN = os.path.join(CORPUS_ROOT, "实施层面详细规划编制工作流程")

CORPUS = [
    # ---- 组一：建筑垃圾规划布置（1 份调研报告 + 2 份汇报，含近似重复稿）----
    {"key": "build_report", "group": "建筑垃圾规划布置", "stage": "调研报告",
     "path": BASE_BUILD + "/关于城镇建筑垃圾处理设施规划布局情况的调研报告.docx"},
    {"key": "build_report_v1", "group": "建筑垃圾规划布置", "stage": "汇报(版1)",
     "path": BASE_BUILD + "/关于城镇建筑垃圾处理设施规划布局情况的汇报.docx"},
    {"key": "build_report_v2", "group": "建筑垃圾规划布置", "stage": "汇报(版2)",
     "path": BASE_BUILD + "/关于城镇建筑垃圾处理设施规划布局情况的汇报 (1).docx"},

    # ---- 组二：实施层面详细规划编制工作流程（7 版迭代稿）----
    {"key": "plan_0312_draft", "group": "规划编制流程", "stage": "草稿",
     "path": BASE_PLAN + "/实施层面详细规划编制工作流程（2025.3.12）（草稿）.docx"},
    {"key": "plan_0312_polish", "group": "规划编制流程", "stage": "润色",
     "path": BASE_PLAN + "/实施层面详细规划编制工作流程（2025.3.12）（润色）.docx"},
    {"key": "plan_0312_review", "group": "规划编制流程", "stage": "审核意见",
     "path": BASE_PLAN + "/实施层面详细规划编制工作流程（2025.3.12）（某某某审）.docx"},
    {"key": "plan_0313_draft", "group": "规划编制流程", "stage": "终稿草稿",
     "path": BASE_PLAN + "/实施层面详细规划编制工作流程（2025.3.13）（终稿草稿）.docx"},
    {"key": "plan_0313", "group": "规划编制流程", "stage": "定稿",
     "path": BASE_PLAN + "/实施层面详细规划编制工作流程（2025.3.13）.docx"},
    {"key": "plan_0318", "group": "规划编制流程", "stage": "迭代",
     "path": BASE_PLAN + "/实施层面详细规划编制工作流程（2025.3.18）.docx"},
    {"key": "plan_0320", "group": "规划编制流程", "stage": "终版",
     "path": BASE_PLAN + "/实施层面详细规划编制工作流程（2025.3.20）.docx"},
]


def clean_text(s):
    """归一化单段文本：去首尾空白、压连续空白为单个空格。"""
    s = s.replace("\u3000", " ")          # 全角空格
    s = re.sub(r"[ \t]+", " ", s)
    return s.strip()


def extract_paragraphs(path):
    """读取 docx，返回非空段落文本列表。"""
    doc = Document(path)
    paras = []
    for p in doc.paragraphs:
        t = clean_text(p.text)
        if t:
            paras.append(t)
    # 表格内的文字也常是正文的一部分（公文里多为要求/清单），一并纳入
    for tbl in doc.tables:
        for row in tbl.rows:
            cells = [clean_text(c.text) for c in row.cells]
            line = " | ".join([c for c in cells if c])
            if line:
                paras.append(line)
    return paras


def est_tokens(text):
    """
    粗略估算 token 数（无需 tokenizer 的近似）：
      汉字/CJK 标点按 1 token/字，拉丁字母与数字按 4 字符 1 token。
    对中文公文而言这个近似已足够判断上下文规模量级。
    """
    cjk = len(re.findall(r"[\u3400-\u9fff\uf900-\ufaff\u3000-\u303f\uff00-\uffef]", text))
    other = len(text) - cjk
    return int(cjk + other / 4.0)


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default="corpus/corpus.json", help="语料输出 JSON 路径")
    args = ap.parse_args()

    out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), args.out)
    os.makedirs(os.path.dirname(out_path), exist_ok=True)

    records = []
    print("%-16s %-14s %8s %6s %8s  %s" % ("key", "阶段", "字符数", "段落", "≈tok", "文件"))
    print("-" * 96)
    for item in CORPUS:
        path = item["path"]
        if not os.path.exists(path):
            print("%-16s %-14s  [缺失] %s" % (item["key"], item["stage"], path))
            continue
        paras = extract_paragraphs(path)
        text = "\n".join(paras)
        rec = {
            "key": item["key"],
            "group": item["group"],
            "stage": item["stage"],
            "path": path,
            "chars": len(text),
            "paras": len(paras),
            "est_tokens": est_tokens(text),
            "text": text,
        }
        records.append(rec)
        print("%-16s %-14s %8d %6d %8d  %s" % (
            item["key"], item["stage"], rec["chars"], rec["paras"],
            rec["est_tokens"], os.path.basename(path)))

    if not records:
        print("未提取到任何语料，请检查路径")
        return 1

    payload = {"generated": time.strftime("%Y-%m-%d %H:%M:%S"), "files": records}
    with open(out_path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=1)
    print("-" * 96)
    print("共 %d 份，合计 %d 字符 / ≈%d tokens" % (
        len(records), sum(r["chars"] for r in records),
        sum(r["est_tokens"] for r in records)))
    print("语料写入 %s" % out_path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
